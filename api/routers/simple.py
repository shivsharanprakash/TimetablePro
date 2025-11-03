import csv
import io
from typing import Dict, Any

from fastapi import APIRouter
from fastapi.responses import StreamingResponse
import zipfile
from io import BytesIO

# import algorithm pieces from project root
import os, sys
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(__file__)))
if ROOT not in sys.path: sys.path.insert(0, ROOT)
from utils import set_up
from scheduler import initial_population, evolutionary_algorithm
from config_adapter import build_data_from_config


router = APIRouter()


def _run(cfg: Dict[str, Any], year: str):
    filled = {}
    subjects_order = {}
    groups_empty_space = {}
    teachers_empty_space = {}
    data = build_data_from_config(cfg, year, teachers_empty_space, groups_empty_space, subjects_order)
    matrix, free = set_up(len(data.classrooms))
    initial_population(data, matrix, free, filled, groups_empty_space, teachers_empty_space, subjects_order)
    evolutionary_algorithm(matrix, data, free, filled, groups_empty_space, teachers_empty_space, subjects_order)
    return matrix


@router.post('/generate')
def generate(payload: Dict[str, Any]):
    cfg = payload
    out = {}
    for year in ['SY', 'TY', 'BTech']:
        out[year] = _run(cfg, year)
    return {'status': 'ok', 'matrices': out}


@router.post('/csv')
def csv_download(payload: Dict[str, Any]):
    cfg = payload['config']
    year = payload['year']
    matrix = _run(cfg, year)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['TimeSlot'] + [f'C{i}' for i in range(len(matrix[0]))])
    for i, row in enumerate(matrix):
        writer.writerow([i] + [('-' if v is None else v) for v in row])
    output.seek(0)
    return StreamingResponse(iter([output.getvalue()]), media_type='text/csv',
                             headers={'Content-Disposition': f'attachment; filename="{year}_timetable.csv"'})


def _build_time_labels(start: str, end: str):
    sh, sm = [int(x) for x in start.split(':')]
    eh, em = [int(x) for x in end.split(':')]
    labels = []
    h = sh
    while h < eh:
        n = h + 1
        labels.append(f"{h:02d}:{sm:02d}-{n:02d}:{sm:02d}")
        h = n
    return labels


def _matrix_to_simple_table(matrix, cfg: Dict[str, Any]):
    # Day headers and rows like the UI, without rowspan merges
    days = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']
    labels = _build_time_labels(cfg['timings']['start'], cfg['timings']['end'])
    slots_per_day = len(labels)
    table = []
    header = ['Time'] + days
    table.append(header)
    for t in range(slots_per_day):
        row = [labels[t]]
        for d in range(6):
            if d == 5:
                row.append('-')
                continue
            mat_row = t + d * 12
            cell = '-'
            for j in range(len(matrix[0])):
                val = matrix[mat_row][j]
                if val is not None:
                    cell = 'X'
                    break
            row.append(cell)
        table.append(row)
    return table


def _csv_bytes(table):
    s = io.StringIO()
    w = csv.writer(s)
    for r in table:
        w.writerow(r)
    return s.getvalue().encode('utf-8')


def _docx_bytes(table):
    from docx import Document
    from docx.shared import Inches
    doc = Document()
    t = doc.add_table(rows=0, cols=len(table[0]))
    for r in table:
        row_cells = t.add_row().cells
        for i, val in enumerate(r):
            row_cells[i].text = str(val)
    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


def _xlsx_bytes(table):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for r in table:
        ws.append([str(c) for c in r])
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


@router.post('/download/all.zip')
def download_all_zip(payload: Dict[str, Any]):
    cfg = payload
    mem = BytesIO()
    with zipfile.ZipFile(mem, mode='w', compression=zipfile.ZIP_DEFLATED) as zf:
        for year in ['SY', 'TY', 'BTech']:
            matrix = _run(cfg, year)
            table = _matrix_to_simple_table(matrix, cfg)
            zf.writestr(f'{year}.docx', _docx_bytes(table))
            zf.writestr(f'{year}.xlsx', _xlsx_bytes(table))
    mem.seek(0)
    return StreamingResponse(mem, media_type='application/zip',
                             headers={'Content-Disposition': 'attachment; filename="timetables_all.zip"'})




